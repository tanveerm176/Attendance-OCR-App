"""Apply readable formatting to the reconciliation Excel report."""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet


HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MATCH_FILL = PatternFill(fill_type="solid", fgColor="E2F0D9")
MISMATCH_FILL = PatternFill(fill_type="solid", fgColor="F4CCCC")
REVIEW_FILL = PatternFill(fill_type="solid", fgColor="FCE4D6")
MISSING_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")
THIN_GRAY_SIDE = Side(style="thin", color="D9E1F2")
CELL_BORDER = Border(
    left=THIN_GRAY_SIDE,
    right=THIN_GRAY_SIDE,
    top=THIN_GRAY_SIDE,
    bottom=THIN_GRAY_SIDE,
)

RESULT_FILLS = {
    "Match": MATCH_FILL,
    "Attendance Mismatch": MISMATCH_FILL,
    "Missing from OCR": MISSING_FILL,
    "Missing from DYCD": MISSING_FILL,
    "Unmatched Roster Name": REVIEW_FILL,
}


def _find_header_column(worksheet: Worksheet, header_name: str) -> int | None:
    """Return the column number for a header, or None when the header is absent."""
    for cell in worksheet[1]:
        if cell.value == header_name:
            return cell.column
    return None


def _format_header_row(worksheet: Worksheet) -> None:
    """Style the column headings and keep them visible while scrolling."""
    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = CELL_BORDER

    worksheet.row_dimensions[1].height = 30
    worksheet.freeze_panes = "A2"


def _format_body_cells(worksheet: Worksheet) -> None:
    """Align report values and add light borders around populated cells."""
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.value is not None:
                cell.border = CELL_BORDER


def _format_result_rows(worksheet: Worksheet) -> None:
    """Color comparison rows according to their Result value."""
    result_column = _find_header_column(worksheet, "Result")
    if result_column is None:
        return

    for row_number in range(2, worksheet.max_row + 1):
        result = worksheet.cell(row=row_number, column=result_column).value
        row_fill = RESULT_FILLS.get(result)
        if row_fill is None:
            continue

        for column_number in range(1, worksheet.max_column + 1):
            worksheet.cell(row=row_number, column=column_number).fill = row_fill


def _format_date_column(worksheet: Worksheet) -> None:
    """Display report dates consistently as month/day/year."""
    date_column = _find_header_column(worksheet, "Date")
    if date_column is None:
        return

    for row_number in range(2, worksheet.max_row + 1):
        worksheet.cell(row=row_number, column=date_column).number_format = "MM/DD/YYYY"


def _size_columns(worksheet: Worksheet) -> None:
    """Set widths from cell contents while keeping very long columns manageable."""
    for column_cells in worksheet.columns:
        column_letter = column_cells[0].column_letter
        longest_value = 0

        for cell in column_cells:
            if cell.value is None:
                continue

            cell_length = len(str(cell.value))
            if cell_length > longest_value:
                longest_value = cell_length

        # A maximum width avoids making the whole report awkward to navigate.
        width = min(max(longest_value + 2, 12), 40)
        worksheet.column_dimensions[column_letter].width = width


def format_report_worksheet(worksheet: Worksheet) -> None:
    """Apply common layout, date, and result-based formatting to one report tab."""
    _format_header_row(worksheet)
    _format_body_cells(worksheet)
    _format_result_rows(worksheet)
    _format_date_column(worksheet)
    _size_columns(worksheet)

    # Filters let reviewers quickly focus on one date, result, or student.
    worksheet.auto_filter.ref = worksheet.dimensions


def format_report_workbook(worksheets: dict[str, Worksheet]) -> None:
    """Format each worksheet in the reconciliation report."""
    for worksheet in worksheets.values():
        format_report_worksheet(worksheet)
