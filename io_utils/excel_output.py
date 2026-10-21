import pandas as pd
from pathlib import Path
from pathlib import PurePosixPath, PureWindowsPath
from openpyxl.worksheet.worksheet import Worksheet
from urllib.parse import quote

from io_utils.excel_formatting import format_attendance_worksheet
from openpyxl.styles import Alignment


OUTPUT_COLUMNS = {
    "Date": "Date",
    "ocr_raw": "OCR Raw",
    "attendance": "Attendance Status",
    "cleaned_name": "OCR Cleaned",
    "matched_name": "Matched Name",
    "matched_score": "Matched Score",
    "flag": "Flag",
}
SOURCE_PDF_PATH_COLUMN = "source_pdf_path"
SOURCE_PDF_PATH_HEADER = "Source PDF Path"


def export_attendance(df: pd.DataFrame, output_path: Path) -> None:
    if "Date" not in df.columns:
        raise ValueError("Attendance data must contain a 'Date' column")
    if df.empty:
        raise ValueError("Cannot export an empty attendance DataFrame")

    missing_columns = set(OUTPUT_COLUMNS).difference(df.columns)
    if missing_columns:
        raise ValueError(
            f"Attendance data is missing required columns: {sorted(missing_columns)}"
        )

    output_columns = dict(OUTPUT_COLUMNS)
    if SOURCE_PDF_PATH_COLUMN in df.columns:
        output_columns[SOURCE_PDF_PATH_COLUMN] = SOURCE_PDF_PATH_HEADER

    attendance = df[list(output_columns)].rename(columns=output_columns).copy()
    attendance["Date"] = pd.to_datetime(attendance["Date"], errors="raise").dt.normalize()
    if attendance["Date"].isna().any():
        raise ValueError("Attendance data contains rows without a date")

    with pd.ExcelWriter(
        output_path,
        engine="openpyxl",
        date_format="MM/DD/YYYY",
        datetime_format="MM/DD/YYYY",
    ) as writer:
        for sheet_date, date_df in attendance.groupby("Date", sort=True):
            # Excel worksheet names cannot contain "/", so use hyphens here.
            sheet_name = sheet_date.strftime("%m-%d-%Y")
            date_df.to_excel(writer, sheet_name=sheet_name, index=False)

            date_column = date_df.columns.get_loc("Date") + 1
            worksheet = writer.sheets[sheet_name]
            for row in worksheet.iter_rows(
                min_row=2,
                min_col=date_column,
                max_col=date_column,
                max_row=worksheet.max_row,
            ):
                row[0].number_format = "MM/DD/YYYY"
            format_attendance_worksheet(worksheet)
            if SOURCE_PDF_PATH_HEADER in date_df.columns:
                _merge_source_pdf_paths(worksheet)

        worksheets = list(writer.sheets.values())
        column_count = max(worksheet.max_column for worksheet in worksheets)
        for column_index in range(1, column_count + 1):
            column_letter = worksheets[0].cell(
                row=1, column=column_index
            ).column_letter
            shared_width = max(
                worksheet.column_dimensions[column_letter].width
                for worksheet in worksheets
            )
            for worksheet in worksheets:
                worksheet.column_dimensions[column_letter].width = shared_width


def _merge_source_pdf_paths(worksheet: Worksheet) -> None:
    source_column = next(
        cell.column
        for cell in worksheet[1]
        if cell.value == SOURCE_PDF_PATH_HEADER
    )
    source_column_letter = worksheet.cell(row=1, column=source_column).column_letter
    worksheet.column_dimensions[source_column_letter].width = min(
        worksheet.column_dimensions[source_column_letter].width, 60
    )

    row = 2
    while row <= worksheet.max_row:
        source_path = worksheet.cell(row=row, column=source_column).value
        if not source_path:
            row += 1
            continue

        end_row = row
        while (
            end_row < worksheet.max_row
            and worksheet.cell(row=end_row + 1, column=source_column).value
            == source_path
        ):
            end_row += 1

        source_name = _source_pdf_name(source_path)
        source_uri = _source_pdf_uri(source_path)
        if end_row > row:
            worksheet.merge_cells(
                start_row=row,
                start_column=source_column,
                end_row=end_row,
                end_column=source_column,
            )
        source_cell = worksheet.cell(row=row, column=source_column)
        source_cell.value = source_name
        source_cell.hyperlink = source_uri
        source_cell.style = "Hyperlink"
        source_cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
            wrap_text=True,
        )
        row = end_row + 1


def _source_pdf_name(source_path: str) -> str:
    if PureWindowsPath(source_path).drive or "\\" in source_path:
        return PureWindowsPath(source_path).name
    return PurePosixPath(source_path).name


def _source_pdf_uri(source_path: str) -> str:
    windows_path = PureWindowsPath(source_path)
    if windows_path.drive or "\\" in source_path:
        normalized_path = str(windows_path).replace("\\", "/")
        if normalized_path.startswith("//"):
            return f"file:{quote(normalized_path, safe='/:')}"
        return f"file:///{quote(normalized_path, safe='/:')}"

    resolved_path = Path(source_path).resolve().as_posix()
    return f"file://{quote(resolved_path, safe='/')}"
