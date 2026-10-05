from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet


def format_attendance_worksheet(worksheet: Worksheet) -> None:
    """Size and style attendance cells, including row highlights by status."""
    centered = Alignment(horizontal="center", vertical="center")
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    present_fill = PatternFill(fill_type="solid", fgColor="8DB4E2")
    absent_fill = PatternFill(fill_type="solid", fgColor="DA9694")
    review_fill = PatternFill(fill_type="solid", fgColor="FFC000")
    thin_side = Side(style="thin", color="000000")
    thin_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    status_column = next(
        (
            cell.column
            for cell in worksheet[1]
            if cell.value == "attendance"
        ),
        None,
    )
    flag_column = next(
        (cell.column for cell in worksheet[1] if cell.value == "flag"),
        None,
    )

    # Let each column grow to fit its longest value, within Excel's width limit.
    for column in worksheet.columns:
        column_letter = column[0].column_letter
        max_length = max(
            (len(str(cell.value)) for cell in column if cell.value is not None),
            default=0,
        )
        worksheet.column_dimensions[column_letter].width = min(max_length + 2, 255)

    for row in worksheet.iter_rows():
        row_fill = None
        if row[0].row > 1:
            status = (
                str(worksheet.cell(row=row[0].row, column=status_column).value)
                if status_column is not None
                else ""
            )
            flag = (
                str(worksheet.cell(row=row[0].row, column=flag_column).value)
                if flag_column is not None
                else ""
            )
            normalized_flag = flag.casefold().replace("-", " ").replace("_", " ")
            if "low match" in normalized_flag or "manual entry" in normalized_flag:
                row_fill = review_fill
            elif status.casefold() == "present":
                row_fill = present_fill
            elif status.casefold() == "absent":
                row_fill = absent_fill

        for cell in row:
            cell.alignment = centered
            if cell.value is not None:
                cell.border = thin_border
                if row_fill is not None:
                    cell.fill = row_fill
            if cell.row == 1:
                cell.fill = header_fill
                cell.font = header_font
