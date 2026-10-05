from datetime import datetime
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from io_utils.excel_output import export_attendance


def test_export_attendance_creates_date_sheets_and_formats_date_column(tmp_path):
    output_path = tmp_path / "attendance.xlsx"
    dataframe = pd.DataFrame(
        {
            "Date": [
                datetime(2026, 10, 2),
                datetime(2026, 9, 15),
                datetime(2026, 10, 2, 15, 30),
                datetime(2026, 10, 2),
            ],
            "Student": ["A", "B", "An exceptionally long student name", "Write-in\nname"],
            "attendance": ["Present", "Absent", "Present", "Present"],
            "flag": [
                "Strong Match",
                "Strong Match",
                "Low Match - Needs Review",
                "NEED-MANUAL-ENTRY",
            ],
        }
    )

    export_attendance(dataframe, output_path)

    workbook = load_workbook(output_path, data_only=False)
    assert workbook.sheetnames == ["09-15-2026", "10-02-2026"]
    assert list(workbook["09-15-2026"].values) == [
        ("Date", "Student", "attendance", "flag"),
        (datetime(2026, 9, 15), "B", "Absent", "Strong Match"),
    ]
    assert workbook["09-15-2026"]["A2"].number_format == "MM/DD/YYYY"
    assert workbook["10-02-2026"]["A2"].number_format == "MM/DD/YYYY"
    assert workbook["10-02-2026"]["A3"].number_format == "MM/DD/YYYY"
    worksheet = workbook["10-02-2026"]
    assert worksheet["A1"].alignment.horizontal == "center"
    assert worksheet["A2"].alignment.horizontal == "center"
    assert worksheet["A1"].font.bold is True
    assert worksheet["A1"].fill.fgColor.rgb == "001F4E78"
    assert worksheet["A2"].border.left.style == "thin"
    assert worksheet["B2"].border.bottom.style == "thin"
    assert worksheet["A2"].fill.fgColor.rgb == "008DB4E2"
    assert worksheet["C2"].fill.fgColor.rgb == "008DB4E2"
    assert workbook["10-02-2026"]["A3"].fill.fgColor.rgb == "00FFC000"
    assert workbook["10-02-2026"]["A4"].fill.fgColor.rgb == "00FFC000"
    assert workbook["09-15-2026"]["A2"].fill.fgColor.rgb == "00DA9694"
    assert worksheet.column_dimensions["B"].width >= len("Student") + 2
    assert worksheet.column_dimensions["B"].width == workbook[
        "09-15-2026"
    ].column_dimensions["B"].width
    assert worksheet.column_dimensions["B"].width >= len(
        "An exceptionally long student name"
    ) + 2
    assert worksheet["B4"].alignment.wrap_text is not True
    assert worksheet.row_dimensions[4].height is None


@pytest.mark.parametrize(
    "dataframe, message",
    [
        (pd.DataFrame({"Student": ["A"]}), "must contain a 'Date' column"),
        (pd.DataFrame({"Date": [datetime(2026, 10, 2)], "Student": ["A"]}).iloc[0:0], "empty"),
        (pd.DataFrame({"Date": [None], "Student": ["A"]}), "without a date"),
    ],
)
def test_export_attendance_rejects_invalid_data(tmp_path, dataframe, message):
    with pytest.raises(ValueError, match=message):
        export_attendance(dataframe, tmp_path / "attendance.xlsx")
