from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from io_utils.excel_output import export_attendance
from ocr_pipeline.batch_cli import process_pdf_folder


class FakePipeline:
    def run(self, pdf_path: Path) -> pd.DataFrame:
        return pd.DataFrame(
            {
                "Date": [datetime(2026, 10, 21), datetime(2026, 10, 21)],
                "ocr_raw": [pdf_path.stem, f"{pdf_path.stem} Student"],
                "attendance": ["Present", "Absent"],
                "cleaned_name": [pdf_path.stem, f"{pdf_path.stem} Student"],
                "matched_name": [pdf_path.stem, f"{pdf_path.stem} Student"],
                "matched_score": [100.0, 100.0],
                "flag": ["Strong Match", "Strong Match"],
                "_review_crop": [object(), object()],
            }
        )


def test_batch_processes_pdfs_and_records_source_paths_without_review_crops(tmp_path):
    (tmp_path / "b.pdf").touch()
    (tmp_path / "A.PDF").touch()

    result = process_pdf_folder(
        tmp_path,
        FakePipeline(),
        source_path_root=r"C:\Attendance\Input",
    )

    assert result["ocr_raw"].tolist() == ["A", "A Student", "b", "b Student"]
    assert result["source_pdf_path"].tolist() == [
        r"C:\Attendance\Input\A.PDF",
        r"C:\Attendance\Input\A.PDF",
        r"C:\Attendance\Input\b.pdf",
        r"C:\Attendance\Input\b.pdf",
    ]
    assert "_review_crop" not in result.columns


def test_export_attendance_merges_source_path_by_pdf_group(tmp_path):
    output_path = tmp_path / "attendance.xlsx"
    dataframe = pd.DataFrame(
        {
            "Date": [datetime(2026, 10, 21)] * 4 + [datetime(2026, 10, 22)],
            "ocr_raw": ["A", "B", "C", "D", "E"],
            "attendance": ["Present", "Absent", "Present", "Absent", "Present"],
            "cleaned_name": ["A", "B", "C", "D", "E"],
            "matched_name": ["A", "B", "C", "D", "E"],
            "matched_score": [100.0] * 5,
            "flag": ["Strong Match"] * 5,
            "source_pdf_path": [
                r"C:\Attendance\Input\one class.pdf",
                r"C:\Attendance\Input\one class.pdf",
                r"C:\Attendance\Input\two.pdf",
                r"C:\Attendance\Input\two.pdf",
                r"C:\Attendance\Input\two.pdf",
            ],
        }
    )

    export_attendance(dataframe, output_path)

    worksheet = load_workbook(output_path)["10-21-2026"]
    assert worksheet["H1"].value == "Source PDF Path"
    assert worksheet["H2"].value == "one class.pdf"
    assert (
        worksheet["H2"].hyperlink.target
        == "file:///C:/Attendance/Input/one%20class.pdf"
    )
    assert worksheet["H4"].value == "two.pdf"
    assert worksheet["H4"].hyperlink.target == "file:///C:/Attendance/Input/two.pdf"
    assert {str(merged) for merged in worksheet.merged_cells.ranges} >= {
        "H2:H3",
        "H4:H5",
    }
    next_date_sheet = load_workbook(output_path)["10-22-2026"]
    assert next_date_sheet["H1"].value == "Source PDF Path"
    assert next_date_sheet["H2"].value == "two.pdf"
    assert next_date_sheet["H2"].hyperlink.target == "file:///C:/Attendance/Input/two.pdf"
